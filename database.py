#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
database.py — 파나소닉 용접로봇 재고관리 시스템 DB
"""

import sqlite3
import json
from datetime import datetime, timezone, timedelta
KST = timezone(timedelta(hours=9))
from werkzeug.security import generate_password_hash
from pathlib import Path
import openpyxl

DB_PATH = Path(__file__).parent / "inventory.db"


def get_db():
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    try:
        conn.execute("PRAGMA journal_mode = WAL")
    except sqlite3.OperationalError:
        # PythonAnywhere의 NFS와 같이 WAL을 지원하지 않는 파일 시스템을 위한 예외 처리
        try:
            conn.execute("PRAGMA journal_mode = DELETE")
        except sqlite3.OperationalError:
            pass
    return conn


def init_db():
    """DB 스키마 생성 + 마이그레이션 + 초기 데이터"""
    conn = get_db()
    try:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS users (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                username      TEXT    UNIQUE NOT NULL,
                password_hash TEXT    NOT NULL,
                name          TEXT    NOT NULL,
                role          TEXT    NOT NULL DEFAULT 'staff',
                is_active     INTEGER NOT NULL DEFAULT 1,
                created_at    DATETIME DEFAULT (datetime('now', '+9 hours'))
            );

            CREATE TABLE IF NOT EXISTS products (
                id              INTEGER PRIMARY KEY AUTOINCREMENT,
                item_code       TEXT    DEFAULT '',
                name            TEXT    NOT NULL,
                category        TEXT    NOT NULL,
                model_spec      TEXT    DEFAULT '',
                price           INTEGER NOT NULL DEFAULT 0,
                quantity        INTEGER NOT NULL DEFAULT 0,
                unit            TEXT    NOT NULL DEFAULT '개',
                location        TEXT    DEFAULT '',
                supplier        TEXT    DEFAULT '',
                low_stock_alert INTEGER NOT NULL DEFAULT 10,
                created_at      DATETIME DEFAULT (datetime('now', '+9 hours')),
                updated_at      DATETIME DEFAULT (datetime('now', '+9 hours'))
            );

            CREATE TABLE IF NOT EXISTS audit_log (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id     INTEGER,
                user_name   TEXT    NOT NULL,
                action      TEXT    NOT NULL,
                target_type TEXT    NOT NULL DEFAULT 'product',
                target_id   INTEGER,
                target_name TEXT,
                before_data TEXT,
                after_data  TEXT,
                note        TEXT    DEFAULT '',
                ip_address  TEXT    DEFAULT '',
                created_at  DATETIME DEFAULT (datetime('now', '+9 hours'))
            );

            CREATE TABLE IF NOT EXISTS work_logs (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id     INTEGER NOT NULL,
                user_name   TEXT    NOT NULL,
                work_date   TEXT    NOT NULL,
                today_work  TEXT    DEFAULT '',
                tomorrow_work TEXT  DEFAULT '',
                memo        TEXT    DEFAULT '',
                content     TEXT    DEFAULT '',
                created_at  DATETIME DEFAULT (datetime('now', '+9 hours')),
                updated_at  DATETIME DEFAULT (datetime('now', '+9 hours'))
            );

            CREATE TABLE IF NOT EXISTS system_config (
                key   TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS backup_history (
                id         INTEGER PRIMARY KEY AUTOINCREMENT,
                filename   TEXT    NOT NULL,
                filepath   TEXT    NOT NULL,
                filesize   INTEGER NOT NULL DEFAULT 0,
                status     TEXT    NOT NULL,
                message    TEXT    DEFAULT '',
                created_at DATETIME DEFAULT (datetime('now', '+9 hours'))
            );
        """)

        # 컬럼 마이그레이션 (기존 DB에 없을 경우 대비)
        migration_cols = [
            ("item_code",  "TEXT DEFAULT ''"),
            ("model_spec", "TEXT DEFAULT ''"),
            ("location",   "TEXT DEFAULT ''"),
        ]
        for col, coltype in migration_cols:
            try:
                conn.execute(f"ALTER TABLE products ADD COLUMN {col} {coltype}")
                conn.commit()
            except Exception:
                pass

        # work_logs 마이그레이션
        try:
            conn.execute("ALTER TABLE work_logs ADD COLUMN today_work TEXT DEFAULT ''")
            conn.execute("ALTER TABLE work_logs ADD COLUMN tomorrow_work TEXT DEFAULT ''")
            conn.execute("ALTER TABLE work_logs ADD COLUMN memo TEXT DEFAULT ''")
            conn.commit()
        except Exception:
            pass

        # 기본 관리자 계정
        if not conn.execute("SELECT id FROM users WHERE username='admin'").fetchone():
            conn.execute(
                "INSERT INTO users (username, password_hash, name, role) VALUES (?,?,?,?)",
                ('admin', generate_password_hash('admin1234'), '관리자', 'admin')
            )

        # 파나소닉 용접로봇 초기 품목 데이터
        if conn.execute("SELECT COUNT(*) FROM products").fetchone()[0] == 0:
            # (item_code, name, category, model_spec, price, quantity, unit, location, supplier, low_stock_alert)
            samples = [
                # ── 소모품 (Consumables) ──
                ("CT-1200", "컨택팁 φ1.2mm",        "소모품",       "Panasonic YX-1200",   1_500,  200, "개",  "A-01-1", "파나소닉코리아",  50),
                ("CT-1600", "컨택팁 φ1.6mm",        "소모품",       "Panasonic YX-1600",   1_800,  150, "개",  "A-01-2", "파나소닉코리아",  50),
                ("NZ-STD",  "노즐 표준형",            "소모품",       "Panasonic NZ-S350",   8_500,   80, "개",  "A-02-1", "파나소닉코리아",  20),
                ("NZ-SML",  "노즐 소형",              "소모품",       "Panasonic NZ-S250",   7_200,   60, "개",  "A-02-2", "파나소닉코리아",  20),
                ("LN-300",  "라이너 3m",              "소모품",       "Panasonic LN-3000",  15_000,   30, "개",  "A-03-1", "파나소닉코리아",  10),
                ("LN-450",  "라이너 4.5m",            "소모품",       "Panasonic LN-4500",  19_000,   20, "개",  "A-03-2", "파나소닉코리아",   8),
                ("WW-12",   "와이어 스풀 φ1.2 15kg", "소모품",       "고려YM-28 φ1.2",     85_000,   20, "릴",  "B-01-1", "고려용접봉",       5),
                ("WW-16",   "와이어 스풀 φ1.6 15kg", "소모품",       "고려YM-28 φ1.6",     88_000,   15, "릴",  "B-01-2", "고려용접봉",       5),
                ("GD-100",  "가스 디퓨저",            "소모품",       "Panasonic GD-100",   12_000,   50, "개",  "A-04-1", "파나소닉코리아",  15),
                ("TH-350",  "용접 토치 350A",         "소모품",       "Panasonic TR-350",  180_000,   10, "개",  "A-05-1", "파나소닉코리아",   3),
                ("TH-500",  "용접 토치 500A",         "소모품",       "Panasonic TR-500",  250_000,    5, "개",  "A-05-2", "파나소닉코리아",   2),

                # ── 주요 모듈 (Main Modules) ──
                ("RC-G3",   "로봇 컨트롤러 G3",       "주요 모듈",   "Panasonic TAWERS G3", 8_500_000, 3, "대", "C-01-1", "파나소닉코리아",   1),
                ("RC-G2",   "로봇 컨트롤러 G2",       "주요 모듈",   "Panasonic TAWERS G2", 6_800_000, 2, "대", "C-01-2", "파나소닉코리아",   1),
                ("TP-DTPS", "티치 펜던트",            "주요 모듈",   "Panasonic DTPS",      1_200_000, 5, "대", "C-02-1", "파나소닉코리아",   1),
                ("SM-082",  "서보 모터 A1 (1축)",     "주요 모듈",   "Panasonic MSMD082",   650_000,   8, "개", "C-03-1", "파나소닉코리아",   2),
                ("SM-152",  "서보 모터 A2 (2축)",     "주요 모듈",   "Panasonic MSMD152",   720_000,   6, "개", "C-03-2", "파나소닉코리아",   2),
                ("DU-540",  "드라이브 유닛",          "주요 모듈",   "Panasonic MDDDT5540", 480_000,   6, "개", "C-04-1", "파나소닉코리아",   2),
                ("CS-5M",   "케이블 세트 5m",         "주요 모듈",   "Panasonic CS-5M",     320_000,  10, "세트","C-05-1","파나소닉코리아",   3),
                ("CS-10M",  "케이블 세트 10m",        "주요 모듈",   "Panasonic CS-10M",    420_000,   6, "세트","C-05-2","파나소닉코리아",   2),

                # ── 전기/전자 부품 (Electrical/Electronic) ──
                ("AS-S1",   "아크 센서",              "전기/전자 부품","Panasonic ARC-S1",   380_000,   4, "개", "D-01-1", "파나소닉코리아",   1),
                ("PS-S2",   "포지션 센서",            "전기/전자 부품","Panasonic POS-S2",   220_000,   6, "개", "D-02-1", "파나소닉코리아",   2),
                ("PC-G3",   "PCB 컨트롤러 보드",      "전기/전자 부품","Panasonic CB-G3",    920_000,   5, "개", "D-03-1", "파나소닉코리아",   1),
                ("PW-24V",  "전원 공급 장치 24V",     "전기/전자 부품","Panasonic PSU-24V",  180_000,   8, "개", "D-04-1", "파나소닉코리아",   2),
                ("PW-48V",  "전원 공급 장치 48V",     "전기/전자 부품","Panasonic PSU-48V",  230_000,   5, "개", "D-04-2", "파나소닉코리아",   2),
                ("RL-24V",  "릴레이 24VDC",           "전기/전자 부품","Omron MY4N-24V",       8_500,  30, "개", "D-05-1", "오므론",          10),
                ("SW-ES",   "비상정지 스위치",        "전기/전자 부품","Idec AB6M-M2",         25_000,  20, "개", "D-06-1", "아이덱",           5),
                ("CN-396",  "커넥터 세트 3.96mm",     "전기/전자 부품","Molex 39-30-1040",     3_500,  50, "세트","D-07-1","몰렉스",          20),
                ("FU-10A",  "퓨즈 10A",               "전기/전자 부품","Bussmann GDC-10A",     1_200, 100, "개", "D-08-1", "버스만",          30),

                # ── 유지보수 도구 및 기타 ──
                ("IK-PM",   "정기점검 키트",          "유지보수/기타","Panasonic PM-KIT",    150_000,  10, "세트","E-01-1","파나소닉코리아",   2),
                ("LG-MK",   "윤활 그리스 400g",       "유지보수/기타","Molykote BR-2Plus",    35_000,  20, "통", "E-02-1", "몰리코트",         5),
                ("CU-FLT",  "냉각장치 필터",          "유지보수/기타","Panasonic CF-G3",      45_000,  15, "개", "E-03-1", "파나소닉코리아",   3),
                ("PC-WLD",  "보호 커버 (용접부)",     "유지보수/기타","범용 HT-Cover",         18_000,  25, "개", "E-04-1", "범용",             5),
                ("SG-HT",   "용접용 내열 안전장갑",   "유지보수/기타","범용 500도 내열",       12_000,  30, "켤레","E-05-1","범용",            10),
                ("SG-FS",   "용접 차광 안면보호대",   "유지보수/기타","범용 자동차광 #9~13",   55_000,  10, "개", "E-05-2", "범용",             3),
                ("TL-SET",  "공구 세트 (정비용)",     "유지보수/기타","범용 17종 세트",        85_000,   5, "세트","E-06-1","범용",             2),
            ]
            conn.executemany(
                """INSERT INTO products
                   (item_code, name, category, model_spec, price, quantity,
                    unit, location, supplier, low_stock_alert)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                samples
            )

        conn.commit()
    finally:
        conn.close()


# ─── Users ────────────────────────────────────
def get_user_by_username(username):
    conn = get_db()
    try:
        row = conn.execute(
            "SELECT * FROM users WHERE username=?", (username,)
        ).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def get_user_by_id(uid):
    conn = get_db()
    try:
        row = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def get_all_users():
    conn = get_db()
    try:
        rows = conn.execute(
            "SELECT * FROM users ORDER BY role DESC, name"
        ).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def create_user(username, password, name, role='staff'):
    conn = get_db()
    try:
        conn.execute(
            "INSERT INTO users (username, password_hash, name, role, created_at) VALUES (?,?,?,?, datetime('now', '+9 hours'))",
            (username, generate_password_hash(password), name, role)
        )
        conn.commit()
        return True
    except sqlite3.IntegrityError:
        return False
    finally:
        conn.close()


def update_user(uid, **kwargs):
    conn = get_db()
    try:
        if 'password' in kwargs:
            kwargs['password_hash'] = generate_password_hash(kwargs.pop('password'))
        fields = ', '.join(f"{k}=?" for k in kwargs)
        values = list(kwargs.values()) + [uid]
        conn.execute(f"UPDATE users SET {fields} WHERE id=?", values)
        conn.commit()
    finally:
        conn.close()


def toggle_user_active(uid):
    conn = get_db()
    try:
        conn.execute(
            "UPDATE users SET is_active=CASE WHEN is_active=1 THEN 0 ELSE 1 END WHERE id=?",
            (uid,)
        )
        conn.commit()
    finally:
        conn.close()


# ─── Products ─────────────────────────────────
def get_all_products(search='', category=''):
    conn = get_db()
    try:
        query = "SELECT * FROM products WHERE 1=1"
        params = []
        if search:
            query += """
              AND (name LIKE ? OR item_code LIKE ? OR model_spec LIKE ?
                   OR supplier LIKE ? OR location LIKE ? OR category LIKE ?)"""
            p = f'%{search}%'
            params += [p, p, p, p, p, p]
        if category and category != '전체':
            query += " AND category=?"
            params.append(category)
        query += " ORDER BY category, item_code"
        rows = conn.execute(query, params).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def get_product(pid):
    conn = get_db()
    try:
        row = conn.execute("SELECT * FROM products WHERE id=?", (pid,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def add_product(item_code, name, category, model_spec, price, quantity,
                unit, location, supplier, low_stock_alert):
    conn = get_db()
    try:
        cur = conn.execute(
            """INSERT INTO products
               (item_code, name, category, model_spec, price, quantity,
                unit, location, supplier, low_stock_alert, created_at, updated_at)
               VALUES (?,?,?,?,?,?,?,?,?,?, datetime('now', '+9 hours'), datetime('now', '+9 hours'))""",
            (item_code, name, category, model_spec, int(price), int(quantity),
             unit, location, supplier, int(low_stock_alert))
        )
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def update_product(pid, **kwargs):
    conn = get_db()
    try:
        kwargs['updated_at'] = datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S")
        fields = ', '.join(f"{k}=?" for k in kwargs)
        values = list(kwargs.values()) + [pid]
        conn.execute(f"UPDATE products SET {fields} WHERE id=?", values)
        conn.commit()
    finally:
        conn.close()


def delete_product(pid):
    conn = get_db()
    try:
        conn.execute("DELETE FROM products WHERE id=?", (pid,))
        conn.commit()
    finally:
        conn.close()


def stock_in(pid, qty):
    conn = get_db()
    try:
        conn.execute(
            "UPDATE products SET quantity=quantity+?, updated_at=datetime('now', '+9 hours') WHERE id=?",
            (int(qty), pid)
        )
        conn.commit()
    finally:
        conn.close()


def stock_out(pid, qty):
    conn = get_db()
    try:
        row = conn.execute("SELECT quantity FROM products WHERE id=?", (pid,)).fetchone()
        if not row or row['quantity'] < int(qty):
            return False
        conn.execute(
            "UPDATE products SET quantity=quantity-?, updated_at=datetime('now', '+9 hours') WHERE id=?",
            (int(qty), pid)
        )
        conn.commit()
        return True
    finally:
        conn.close()


def get_categories():
    conn = get_db()
    try:
        rows = conn.execute(
            "SELECT DISTINCT category FROM products ORDER BY category"
        ).fetchall()
        return [r['category'] for r in rows]
    finally:
        conn.close()


def clear_all_products():
    """Clear all products from the products table and reset sqlite sequence"""
    conn = get_db()
    try:
        conn.execute("DELETE FROM products")
        conn.execute("DELETE FROM sqlite_sequence WHERE name='products'")
        conn.commit()
    finally:
        conn.close()


def get_product_by_item_code(item_code):
    """Retrieve a product dict by item code"""
    conn = get_db()
    try:
        row = conn.execute("SELECT * FROM products WHERE item_code=?", (item_code,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def get_category_by_name(name):
    if not name:
        return "유지보수/기타"
    name_upper = str(name).upper()
    if any(k in name_upper for k in ['TIP', 'NOZZLE', 'LINER', 'WIRE', 'DIFFUSER', 'TORCH', 'SPRING', 'PIN', 'BOLT', 'NUT', 'ROLLER', 'GEAR', 'O-RING', 'PACKING', 'GUIDE', 'ORING', 'WASHER', 'COLLET', 'TUBE', 'PAD', 'BELT', 'HOLDER', 'INSULATOR', 'BUSH', 'GASKET', 'SEAL', 'CLAMP']):
        return "소모품"
    elif any(k in name_upper for k in ['MOTOR', 'DRIVE', 'PENDANT', 'CONTROLLER', 'GEARBOX', 'REDUCTION', 'HARMONIC', 'UNIT', 'MODULE', 'DISPLAY', 'LCD', 'TRANSFORMER']):
        return "주요 모듈"
    elif any(k in name_upper for k in ['BOARD', 'CABLE', 'SENSOR', 'SWITCH', 'RELAY', 'POWER', 'DIODE', 'FUSE', 'CONNECTOR', 'PLUG', 'RECEPTACLE', 'BREAKER', 'VALVE', 'SOLENOID', 'FAN', 'HEATER', 'ENCODER', 'PCB', 'SWTICH']):
        return "전기/전자 부품"
    return "유지보수/기타"


def import_excel_to_db(file_stream, user_id=0, user_name='시스템'):
    """Parse Excel file from file_stream and load it into the products table, clearing existing data first"""
    # Load optional prices mapping from 부품가격표.xlsx if it exists
    price_map = {}
    price_file = Path(__file__).parent / "부품가격표.xlsx"
    if price_file.exists():
        try:
            price_wb = openpyxl.load_workbook(str(price_file), data_only=True)
            price_sheet = price_wb.active
            for idx, row in enumerate(price_sheet.iter_rows(values_only=True)):
                if idx <= 1:
                    continue
                code = str(row[0]).strip() if row[0] is not None else ""
                val = row[2]
                if code and val is not None:
                    try:
                        price_map[code] = int(val)
                    except (ValueError, TypeError):
                        pass
        except Exception as e:
            pass

    wb = openpyxl.load_workbook(file_stream, data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    data_rows = rows[2:]  # Skip top empty row and header row
    
    conn = get_db()
    try:
        # Clear existing data
        conn.execute("DELETE FROM products")
        conn.execute("DELETE FROM sqlite_sequence WHERE name='products'")
        
        inserted_count = 0
        skipped_count = 0
        
        for row in data_rows:
            item_code = str(row[0]).strip() if row[0] is not None else ""
            name = str(row[1]).strip() if row[1] is not None else ""
            
            # If both are empty, skip
            if not item_code and not name:
                skipped_count += 1
                continue
                
            qty = 0
            if row[2] is not None:
                try:
                    qty = int(row[2])
                except (ValueError, TypeError):
                    pass
                    
            spec_parts = []
            if row[3] is not None:
                spec_parts.append(str(row[3]).strip())
            if row[4] is not None:
                spec_parts.append(f"{row[4]}개")
            model_spec = " / ".join(spec_parts)
            
            category = get_category_by_name(name)
            location = str(row[5]).strip() if row[5] is not None else ""
            
            # Defaults
            price = price_map.get(item_code, 0)
            unit = "개"
            supplier = ""
            low_stock_alert = 10
            
            conn.execute(
                """INSERT INTO products 
                   (item_code, name, category, model_spec, price, quantity, unit, location, supplier, low_stock_alert, created_at, updated_at) 
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now', '+9 hours'), datetime('now', '+9 hours'))""",
                (item_code, name, category, model_spec, price, qty, unit, location, supplier, low_stock_alert)
            )
            inserted_count += 1
            
        # Add audit log
        conn.execute(
            """INSERT INTO audit_log 
               (user_id, user_name, action, target_type, note, created_at) 
               VALUES (?, ?, ?, ?, ?, datetime('now', '+9 hours'))""",
            (user_id, user_name, "BULK_IMPORT", "product", f"엑셀 파일 일괄 등록 완료 (신규 {inserted_count}건 등록, {skipped_count}건 스킵)")
        )
        conn.commit()
        return inserted_count, skipped_count
    except Exception as e:
        conn.rollback()
        raise e
    finally:
        conn.close()


def update_prices_from_excel(file_stream, user_id=0, user_name='시스템'):
    """Parse Excel price list file from file_stream and update prices of existing items matching by item_code"""
    wb = openpyxl.load_workbook(file_stream, data_only=True)
    sheet = wb.active
    
    excel_prices = {}
    for idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if idx <= 1: # Skip header/empty rows
            continue
        item_code = str(row[0]).strip() if row[0] is not None else ""
        price = row[2]
        if item_code and price is not None:
            try:
                excel_prices[item_code] = int(price)
            except (ValueError, TypeError):
                pass
                
    conn = get_db()
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT id, item_code, name, price FROM products")
        products = cursor.fetchall()
        
        updated_count = 0
        no_change_count = 0
        no_match_count = 0
        
        for p in products:
            pid = p['id']
            item_code = p['item_code'].strip() if p['item_code'] else ""
            old_price = p['price']
            
            if not item_code:
                no_match_count += 1
                continue
                
            if item_code in excel_prices:
                new_price = excel_prices[item_code]
                if old_price != new_price:
                    conn.execute(
                        "UPDATE products SET price=?, updated_at=datetime('now', '+9 hours') WHERE id=?",
                        (new_price, pid)
                    )
                    updated_count += 1
                else:
                    no_change_count += 1
            else:
                no_match_count += 1
                
        if updated_count > 0:
            conn.execute(
                """INSERT INTO audit_log 
                   (user_id, user_name, action, target_type, note, created_at) 
                   VALUES (?, ?, ?, ?, ?, datetime('now', '+9 hours'))""",
                (user_id, user_name, "EDIT", "product", f"부품가격표 업로드를 통한 단가 일괄 업데이트 완료 ({updated_count}건 수정, 변경없음 {no_change_count}건, 미매칭 {no_match_count}건)")
            )
            conn.commit()
            
        return updated_count, no_change_count, no_match_count
    except Exception as e:
        conn.rollback()
        raise e
    finally:
        conn.close()



# ─── Audit Log ────────────────────────────────
def add_audit_log(user_id, user_name, action, target_type,
                  target_id, target_name,
                  before_data=None, after_data=None,
                  note='', ip_address=''):
    conn = get_db()
    try:
        conn.execute(
            """INSERT INTO audit_log
               (user_id, user_name, action, target_type, target_id, target_name,
                before_data, after_data, note, ip_address, created_at)
               VALUES (?,?,?,?,?,?,?,?,?,?, datetime('now', '+9 hours'))""",
            (
                user_id, user_name, action, target_type, target_id, target_name,
                json.dumps(before_data, ensure_ascii=False) if before_data else None,
                json.dumps(after_data,  ensure_ascii=False) if after_data  else None,
                note, ip_address
            )
        )
        conn.commit()
    finally:
        conn.close()


def get_audit_logs(limit=300, action_filter='', user_filter='', search=''):
    conn = get_db()
    try:
        query = "SELECT * FROM audit_log WHERE 1=1"
        params = []
        if action_filter:
            query += " AND action=?"
            params.append(action_filter)
        if user_filter:
            query += " AND user_name LIKE ?"
            params.append(f'%{user_filter}%')
        if search:
            query += " AND (target_name LIKE ? OR note LIKE ? OR user_name LIKE ?)"
            params += [f'%{search}%', f'%{search}%', f'%{search}%']
        query += " ORDER BY created_at DESC LIMIT ?"
        params.append(limit)
        rows = conn.execute(query, params).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


# ─── Dashboard stats ──────────────────────────
def get_stats():
    conn = get_db()
    try:
        total_products = conn.execute("SELECT COUNT(*) FROM products").fetchone()[0]
        total_value    = conn.execute(
            "SELECT COALESCE(SUM(price*quantity),0) FROM products"
        ).fetchone()[0]
        low_stock = conn.execute(
            "SELECT COUNT(*) FROM products WHERE quantity<=low_stock_alert"
        ).fetchone()[0]
        today = datetime.now(KST).strftime('%Y-%m-%d')
        today_tx = conn.execute(
            "SELECT COUNT(*) FROM audit_log WHERE date(created_at)=? AND action IN ('STOCK_IN','STOCK_OUT')",
            (today,)
        ).fetchone()[0]

        # 카테고리별 품목 수
        cat_rows = conn.execute(
            "SELECT category, COUNT(*) as cnt, SUM(price*quantity) as val FROM products GROUP BY category ORDER BY category"
        ).fetchall()
        categories = [{'category': r['category'], 'count': r['cnt'], 'value': r['val'] or 0}
                      for r in cat_rows]

        return {
            'total_products':     total_products,
            'total_value':        total_value,
            'low_stock':          low_stock,
            'today_transactions': today_tx,
            'categories':         categories,
        }
    finally:
        conn.close()


# ─── Work Logs ────────────────────────────────
def get_work_logs(user_id=None, work_date_start=None, work_date_end=None, search_keyword=None, limit=300):
    conn = get_db()
    try:
        query = "SELECT * FROM work_logs WHERE 1=1"
        params = []
        if user_id is not None:
            query += " AND user_id=?"
            params.append(user_id)
        if work_date_start:
            query += " AND work_date >= ?"
            params.append(work_date_start)
        if work_date_end:
            query += " AND work_date <= ?"
            params.append(work_date_end)
        if search_keyword:
            query += " AND (today_work LIKE ? OR tomorrow_work LIKE ? OR memo LIKE ? OR content LIKE ?)"
            kw = f"%{search_keyword}%"
            params.extend([kw, kw, kw, kw])
        query += " ORDER BY work_date DESC, created_at DESC LIMIT ?"
        params.append(limit)
        rows = conn.execute(query, params).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def get_work_log(log_id):
    conn = get_db()
    try:
        row = conn.execute("SELECT * FROM work_logs WHERE id=?", (log_id,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def add_work_log(user_id, user_name, work_date, today_work, tomorrow_work, memo):
    conn = get_db()
    try:
        cur = conn.execute(
            """INSERT INTO work_logs (user_id, user_name, work_date, today_work, tomorrow_work, memo, content, created_at, updated_at) 
               VALUES (?, ?, ?, ?, ?, ?, '', datetime('now', '+9 hours'), datetime('now', '+9 hours'))""",
            (user_id, user_name, work_date, today_work, tomorrow_work, memo)
        )
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def update_work_log(log_id, today_work, tomorrow_work, memo):
    conn = get_db()
    try:
        conn.execute(
            "UPDATE work_logs SET today_work=?, tomorrow_work=?, memo=?, updated_at=datetime('now', '+9 hours') WHERE id=?",
            (today_work, tomorrow_work, memo, log_id)
        )
        conn.commit()
    finally:
        conn.close()


def delete_work_log(log_id):
    conn = get_db()
    try:
        conn.execute("DELETE FROM work_logs WHERE id=?", (log_id,))
        conn.commit()
    finally:
        conn.close()


# ─── System Config & Backup History ─────────────
DEFAULT_CONFIG = {
    'email_provider': 'smtp',  # 'smtp', 'resend', 'brevo'
    'api_key': '',
    'smtp_host': 'smtp.naver.com',
    'smtp_port': '587',
    'smtp_user': '',
    'smtp_pass': '',
    'smtp_sender': '',
    'smtp_receiver': '',
    'smtp_use_tls': '1',
    'auto_backup_enabled': '1',
    'auto_backup_time': '00:00',
}


def get_system_config():
    conn = get_db()
    try:
        rows = conn.execute("SELECT key, value FROM system_config").fetchall()
        cfg = dict(DEFAULT_CONFIG)
        for r in rows:
            cfg[r['key']] = r['value']
        return cfg
    finally:
        conn.close()


def set_system_config(config_dict):
    conn = get_db()
    try:
        for k, v in config_dict.items():
            conn.execute(
                "INSERT INTO system_config (key, value) VALUES (?, ?) "
                "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                (str(k), str(v))
            )
        conn.commit()
    finally:
        conn.close()


def add_backup_history(filename, filepath, filesize, status, message=''):
    conn = get_db()
    try:
        conn.execute(
            """INSERT INTO backup_history (filename, filepath, filesize, status, message, created_at)
               VALUES (?, ?, ?, ?, ?, datetime('now', '+9 hours'))""",
            (filename, filepath, filesize, status, message)
        )
        conn.commit()
    finally:
        conn.close()


def get_backup_history(limit=50):
    conn = get_db()
    try:
        rows = conn.execute(
            "SELECT * FROM backup_history ORDER BY id DESC LIMIT ?", (limit,)
        ).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


