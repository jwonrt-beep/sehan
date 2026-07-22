#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
import_excel.py - 엑셀 파일 부품코드및재고량.xlsx 데이터를 DB에 입력하는 스크립트
"""

import sys
import os
import openpyxl
import sqlite3
from pathlib import Path

# Workspace 경로 설정
WORKSPACE_DIR = Path(__file__).parent
DB_PATH = WORKSPACE_DIR / "inventory.db"
EXCEL_PATH = WORKSPACE_DIR / "부품코드및재고량.xlsx"

def get_category(name):
    if not name:
        return "유지보수/기타"
    name_upper = str(name).upper()
    if any(k in name_upper for k in ['TIP', 'NOZZLE', 'LINER', 'WIRE', 'DIFFUSER', 'TORCH', 'SPRING', 'PIN', 'BOLT', 'NUT', 'ROLLER', 'GEAR', 'O-RING', 'PACKING', 'GUIDE', 'ORING', 'WASHER', 'COLLET', 'TUBE', 'PAD', 'BELT', 'HOLDER', 'INSULATOR', 'BUSH', 'GASKET', 'SEAL', 'CLAMP']):
        return "소모품"
    elif any(k in name_upper for k in ['MOTOR', 'DRIVE', 'PENDANT', 'CONTROLLER', 'GEARBOX', 'REDUCTION', 'HARMONIC', 'UNIT', 'MODULE', 'DISPLAY', 'LCD', 'TRANSFORMER']):
        return "주요 모듈"
    elif any(k in name_upper for k in ['BOARD', 'CABLE', 'SENSOR', 'SWITCH', 'RELAY', 'POWER', 'DIODE', 'FUSE', 'CONNECTOR', 'PLUG', 'RECEPTACLE', 'BREAKER', 'VALVE', 'SOLENOID', 'FAN', 'HEATER', 'ENCODER', 'PCB', 'SWTICH']):
        return "전기/전자 부품"
    return "유지보수/기타"

def main():
    if not EXCEL_PATH.exists():
        print(f"[오류] 엑셀 파일을 찾을 수 없습니다: {EXCEL_PATH}")
        sys.exit(1)
        
    # 가격 정보 로드 (부품가격표.xlsx가 존재하는 경우)
    price_map = {}
    price_file = WORKSPACE_DIR / "부품가격표.xlsx"
    if price_file.exists():
        print(f"[*] 가격표 파일 발견: {price_file.name}. 단가 데이터를 매핑합니다...")
        try:
            price_wb = openpyxl.load_workbook(price_file, data_only=True)
            price_sheet = price_wb.active
            for idx, row in enumerate(price_sheet.iter_rows(values_only=True)):
                if idx <= 1:
                    continue
                code = str(row[0]).strip() if row[0] is not None else ""
                val = row[2]
                if code and val is not None:
                    try:
                        price_map[code] = int(val)
                    except (ValueError, TypeError):
                        pass
            print(f"[+] 가격표 매핑 완료: 총 {len(price_map)}개 항목 로드.")
        except Exception as e:
            print(f"[경고] 가격표 파일을 읽는 데 실패했습니다: {e}")

    print(f"[*] 엑셀 파일 로딩 중: {EXCEL_PATH.name}...")
    try:
        wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
        sheet = wb.active
        print(f"[*] 활성 시트: {sheet.title}")
    except Exception as e:
        print(f"[오류] 엑셀 파일을 읽는 도중 오류가 발생했습니다: {e}")
        sys.exit(1)

    rows = list(sheet.iter_rows(values_only=True))
    # 1번째 줄 비어있고, 2번째 줄 헤더, 3번째 줄부터 데이터
    data_rows = rows[2:]
    print(f"[*] 총 데이터 행 수: {len(data_rows)}개")

    # DB 연결 및 비우기
    if not DB_PATH.exists():
        print(f"[*] DB 파일이 존재하지 않아 새로 생성합니다: {DB_PATH}")
    
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    try:
        # 기존 데이터 삭제
        print("[*] 기존 재고 품목 데이터를 모두 삭제합니다...")
        cursor.execute("DELETE FROM products")
        cursor.execute("DELETE FROM sqlite_sequence WHERE name='products'")
        
        inserted_count = 0
        skipped_count = 0

        # 데이터 파싱 및 삽입
        print("[*] 엑셀 데이터 파싱 및 DB 삽입 진행 중...")
        for idx, row in enumerate(data_rows):
            # 컬럼 매핑:
            # 0: Part Number for Order (item_code)
            # 1: Description (name)
            # 2: 수량 (quantity)
            # 3: Note / Spec (예: 중고3, 추가5 등)
            # 4: Used Quantity (예: 2, 4 등)
            # 5: Location (예: 11,44)
            item_code = str(row[0]).strip() if row[0] is not None else ""
            name = str(row[1]).strip() if row[1] is not None else ""
            
            # 품명과 품목코드가 모두 없으면 스킵
            if not item_code and not name:
                skipped_count += 1
                continue

            # 수량 파싱
            qty = 0
            if row[2] is not None:
                try:
                    qty = int(row[2])
                except (ValueError, TypeError):
                    pass

            # 규격 및 중고 비고 조합
            spec_parts = []
            if row[3] is not None:
                spec_parts.append(str(row[3]).strip())
            if row[4] is not None:
                spec_parts.append(f"{row[4]}개")
            model_spec = " / ".join(spec_parts)

            # 카테고리 자동 판별
            category = get_category(name)

            # 위치
            location = str(row[5]).strip() if row[5] is not None else ""

            # 기본값 설정
            price = price_map.get(item_code, 0)
            unit = "개"
            supplier = ""
            low_stock_alert = 10

            cursor.execute(
                """INSERT INTO products 
                   (item_code, name, category, model_spec, price, quantity, unit, location, supplier, low_stock_alert, created_at, updated_at) 
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now', '+9 hours'), datetime('now', '+9 hours'))""",
                (item_code, name, category, model_spec, price, qty, unit, location, supplier, low_stock_alert)
            )
            inserted_count += 1

        # 감사 로그 작성
        cursor.execute(
            """INSERT INTO audit_log 
               (user_id, user_name, action, target_type, note, created_at) 
               VALUES (?, ?, ?, ?, ?, datetime('now', '+9 hours'))""",
            (0, "시스템(CLI)", "BULK_IMPORT", "product", f"엑셀 파일 일괄 등록 완료 (신규 {inserted_count}건 등록, {skipped_count}건 스킵)")
        )

        conn.commit()
        print(f"[+] 성공: {inserted_count}개의 품목을 등록했습니다. (스킵된 빈 행: {skipped_count}개)")

    except Exception as e:
        conn.rollback()
        print(f"[오류] 데이터베이스 처리 중 오류가 발생하여 롤백했습니다: {e}")
        sys.exit(1)
    finally:
        conn.close()

if __name__ == "__main__":
    main()
